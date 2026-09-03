#!/usr/bin/env python3
"""Apply deterministic publication metadata and page furniture to a DOCX."""

from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "Psychology Statistics Practice Materials with Jamovi, Volume 2"
SUBTITLE = "Categorical, Rank-Based, Quasi-Experimental, and Single-Case Evidence"
DOI = "10.5281/zenodo.22286929"
RUNNING_TITLE = "Psychology Statistics Practice Materials with Jamovi, Volume 2"
FIXED_TIME = dt.datetime(2026, 9, 3, 0, 0, 0)


def remove_children(element) -> None:
    for child in list(element):
        element.remove(child)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for old in tr_pr.findall(qn("w:tblHeader")):
        tr_pr.remove(old)
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def replace_child(parent, tag: str, value) -> None:
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    parent.append(value)


def set_cell_margins(cell, *, top: int, left: int, bottom: int, right: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, amount in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")


def column_widths(table, total_width: int) -> list[int]:
    maxima: list[int] = []
    for column_index in range(len(table.columns)):
        lengths = [
            max((len(line) for line in cell.text.splitlines()), default=1)
            for cell in (row.cells[column_index] for row in table.rows)
        ]
        maxima.append(max(lengths or [1]))
    weights = [max(2.2, min(9.5, math.sqrt(length))) for length in maxima]
    raw = [total_width * weight / sum(weights) for weight in weights]
    widths = [int(value) for value in raw]
    widths[-1] += total_width - sum(widths)
    return widths


def style_table(table) -> None:
    column_count = len(table.columns)
    total_width = 9120
    widths = column_widths(table, total_width)
    table.style = "Table Grid"
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    replace_child(tbl_pr, "w:tblW", tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    replace_child(tbl_pr, "w:tblInd", tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    replace_child(tbl_pr, "w:tblLayout", layout)

    grid = table._tbl.tblGrid
    remove_children(grid)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    font_size = 8.0 if column_count >= 8 else 9.0 if column_count >= 6 else 9.5
    margin = 55 if column_count >= 8 else 80
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        replace_child(tr_pr, "w:cantSplit", cant_split)
        for column_index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[column_index]))
            tc_w.set(qn("w:type"), "dxa")
            replace_child(tc_pr, "w:tcW", tc_w)
            set_cell_margins(
                cell,
                top=70,
                left=margin,
                bottom=70,
                right=margin,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                replace_child(tc_pr, "w:shd", shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.widow_control = True
                if column_count >= 6 and column_index > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        "Liberation Serif",
                        font_size,
                        bold=True if row_index == 0 else None,
                    )


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, value, end):
        run._r.append(node)
    run.font.name = "Liberation Serif"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)


def set_run_font(run, name: str, size: float, *, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def has_page_break(paragraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def style_cover(document: Document) -> None:
    cover: list = []
    for paragraph in document.paragraphs:
        if has_page_break(paragraph):
            break
        cover.append(paragraph)
    title_count = sum(paragraph.text.strip() == TITLE for paragraph in cover)
    if title_count != 1:
        raise SystemExit(f"DOCX_ERROR: cover title count is {title_count}, expected 1")

    for paragraph in cover:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(9)
        style_name = paragraph.style.name if paragraph.style is not None else ""
        for run in paragraph.runs:
            if style_name == "Title" or paragraph.text.strip() == TITLE:
                set_run_font(run, "Liberation Sans", 24, bold=True, color=RGBColor(32, 55, 72))
            elif style_name == "Subtitle" or paragraph.text.strip() == SUBTITLE:
                set_run_font(run, "Liberation Sans", 14, italic=True, color=RGBColor(65, 80, 92))
            elif style_name == "Author":
                set_run_font(run, "Liberation Serif", 12, bold=True)
            elif style_name == "Date":
                set_run_font(run, "Liberation Serif", 10.5)
            else:
                set_run_font(run, "Liberation Serif", 10.5)
    cover[0].paragraph_format.space_before = Pt(112)


def request_field_update(document: Document) -> None:
    settings = document.settings.element
    old = settings.find(qn("w:updateFields"))
    if old is not None:
        settings.remove(old)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def repair_missing_paragraph_styles(document: Document) -> None:
    """Map Pandoc paragraph styles absent from Volume 1 to BodyText.

    The Volume 1 reference DOCX defines neither FirstParagraph nor Compact.
    LibreOffice otherwise renders FirstParagraph with a literal trailing "X"
    in some linked paragraphs and suppresses Compact table-cell text entirely.
    Search the complete document tree because ``document.paragraphs`` omits
    paragraphs nested inside tables.
    """
    missing_styles = {"FirstParagraph", "Compact"}
    for p_style in document.element.xpath(".//w:pStyle"):
        if p_style.get(qn("w:val")) in missing_styles:
            p_style.set(qn("w:val"), "BodyText")


def normalize_docx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            old = source.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 3, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = old.external_attr or ((0o100644 & 0xFFFF) << 16)
            target.writestr(info, source.read(name))
    os.replace(temporary, path)


parser = argparse.ArgumentParser()
parser.add_argument("docx", type=Path)
args = parser.parse_args()

document = Document(args.docx)
properties = document.core_properties
properties.title = f"{TITLE}: {SUBTITLE}"
properties.subject = f"Version 1.0; CC BY 4.0; DOI {DOI}"
properties.author = "Nicholas Elliott Karlson"
properties.category = "Open educational resource - Lesson"
properties.comments = (
    "Four units, 96 practice problems, 96 fully worked solutions, synthetic data, "
    "and independently verified results. Publisher: NEKpress Research."
)
properties.identifier = DOI
properties.keywords = f"{DOI}; CC BY 4.0; jamovi; psychology statistics; open education"
properties.language = "en-US"
properties.last_modified_by = "NEKpress Research"
properties.created = FIXED_TIME
properties.modified = FIXED_TIME
properties.revision = 1

# Volume 1's DOCX is the Pandoc reference document and carries
# ``evenAndOddHeaders``.  If that switch survives, Word and LibreOffice apply
# our default header and footer to odd pages only, leaving every even page
# without its running title or page number.
settings = document.settings.element
for node in settings.findall(qn("w:evenAndOddHeaders")):
    settings.remove(node)

for section in document.sections:
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    remove_children(paragraph._p)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(RUNNING_TITLE)
    set_run_font(run, "Liberation Sans", 8.5, color=RGBColor(90, 90, 90))

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    for first_paragraph in first_header.paragraphs:
        remove_children(first_paragraph._p)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    remove_children(paragraph._p)
    add_page_field(paragraph)

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    for first_paragraph in first_footer.paragraphs:
        remove_children(first_paragraph._p)

style_cover(document)
repair_missing_paragraph_styles(document)

for paragraph in document.paragraphs:
    if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.widow_control = True

for table in document.tables:
    style_table(table)
    if table.rows:
        set_repeat_table_header(table.rows[0])

request_field_update(document)
document.save(args.docx)
normalize_docx(args.docx)
print("DOCX_POSTPROCESS_OK")
